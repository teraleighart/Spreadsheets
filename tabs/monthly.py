"""
Monthly tab builder — reused for January through December.
Every monthly tab has identical structure; the month name is parameterised.

Fixed row anchors (used by Annual Totals cross-refs):
  MTH_INCOME_TOTAL_ROW  = 13   → F13 = total actual income
  MTH_EXPENSE_TOTAL_ROW = 37   → F37 = total actual expenses
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.utils import get_column_letter

from styling import (
    COLORS, FONT_DISPLAY, FONT_BODY,
    setup_sheet, add_gold_rule,
    add_dropdown, add_data_bars,
)
from constants import (
    EXPENSE_CATS, INCOME_CATS, MONTHS, sidebar,
    MTH_INCOME_TOTAL_ROW, MTH_EXPENSE_TOTAL_ROW,
)

# ── column map ────────────────────────────────────────────────
#   D=4  Category / Source
#   E=5  Budgeted
#   F=6  Actual
#   G=7  Remaining  (Budgeted − Actual)
#   H=8  % Used
#   I=9  Status

# ── row map ───────────────────────────────────────────────────
_INCOME_HDR_ROW  = 7     # section header
_INCOME_COL_ROW  = 8     # column labels
_INCOME_DATA_START = 9   # first income data row
_INCOME_DATA_END   = 12  # last income data row  (4 income sources)
_INCOME_TOTAL_ROW  = MTH_INCOME_TOTAL_ROW   # = 13

_EXPENSE_HDR_ROW    = 16
_EXPENSE_COL_ROW    = 17
_EXPENSE_DATA_START = 18
_EXPENSE_DATA_END   = 36   # 19 expense categories (18–36)
_EXPENSE_TOTAL_ROW  = MTH_EXPENSE_TOTAL_ROW  # = 37

_SUMMARY_HDR_ROW  = 40
_SUMMARY_COL_ROW  = 41
_SUMMARY_DATA_ROW = 42   # Total Income
# 43 = Total Expenses
# 44 = Net Savings
# 45 = Savings Rate


def _W(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _sec(ws, row, col, text, end_col=None, fill_color=None):
    c = ws.cell(row=row, column=col)
    c.value = f"  {text}"
    c.style = 'section_header'
    if fill_color:
        c.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
    if end_col:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=end_col)


def _col_hdr(ws, row, col, labels):
    for i, lbl in enumerate(labels):
        c = ws.cell(row=row, column=col + i)
        c.value = lbl
        c.style = 'column_header'


def _data(ws, row, col, vals, alt=False, fmt=None):
    sn = 'data_alt' if alt else 'data_cell'
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=col + i)
        c.value = v
        c.style = sn
        if isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith('=')):
            c.alignment = Alignment(horizontal='right', vertical='center')
        if fmt and i < len(fmt) and fmt[i]:
            c.number_format = fmt[i]


def _total(ws, row, col, vals):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=col + i)
        c.value = v
        c.style = 'total_row'
        if i == 0:
            c.alignment = Alignment(horizontal='left', vertical='center')
        else:
            c.alignment = Alignment(horizontal='right', vertical='center')
            if isinstance(v, str) and v.startswith('='):
                c.number_format = '"$"#,##0.00'


def _money(ws, row, col, count=1):
    for i in range(count):
        ws.cell(row=row, column=col + i).number_format = '"$"#,##0.00'


def _pct(ws, row, col):
    ws.cell(row=row, column=col).number_format = '0.0%'


def _add_over_budget_cf(ws, start_row, end_row, col_g=7):
    """Red fill when Remaining (col G) < 0 (over budget)."""
    cell_range = f'G{start_row}:G{end_row}'
    over_style = DifferentialStyle(
        fill=PatternFill(start_color=COLORS['danger'], end_color=COLORS['danger'], fill_type='solid'),
        font=Font(name=FONT_BODY, size=10, bold=True, color=COLORS['white']),
    )
    near_style = DifferentialStyle(
        fill=PatternFill(start_color=COLORS['gold'], end_color=COLORS['gold'], fill_type='solid'),
    )
    ok_style = DifferentialStyle(
        fill=PatternFill(start_color=COLORS['success'], end_color=COLORS['success'], fill_type='solid'),
    )
    ws.conditional_formatting.add(cell_range,
        Rule(type='cellIs', operator='lessThan', formula=['0'], dxf=over_style))
    ws.conditional_formatting.add(cell_range,
        Rule(type='cellIs', operator='between', formula=['0', '50'], dxf=near_style))
    ws.conditional_formatting.add(cell_range,
        Rule(type='cellIs', operator='greaterThan', formula=['50'], dxf=ok_style))


# ═══════════════════════════════════════════════════════════════
# MAIN BUILDER
# ═══════════════════════════════════════════════════════════════

def build_monthly_tab(wb, ws, month_name):
    """Build one monthly budget tab for the given month_name (e.g. 'January')."""
    year_cell = "='Setup'!F8"   # year from Setup tab

    setup_sheet(ws, month_name,
        subtitle=f"Budget vs. actual for {month_name}. Spending figures pull automatically from Variable Transactions.",
        active_tab=month_name, sections=sidebar(month_name))

    _W(ws, {'D': 26, 'E': 14, 'F': 14, 'G': 14, 'H': 11, 'I': 16})

    # Row heights for key rows
    for r in [_INCOME_HDR_ROW, _EXPENSE_HDR_ROW, _SUMMARY_HDR_ROW]:
        ws.row_dimensions[r].height = 24
    for r in [_INCOME_COL_ROW, _EXPENSE_COL_ROW, _SUMMARY_COL_ROW]:
        ws.row_dimensions[r].height = 20

    # ── NAV ARROWS ────────────────────────────────────────────
    idx = MONTHS.index(month_name)
    prev_month = MONTHS[idx - 1] if idx > 0 else None
    next_month = MONTHS[idx + 1] if idx < 11 else None

    if prev_month:
        pc = ws.cell(row=2, column=14)
        pc.value = f"← {prev_month}"
        pc.hyperlink = f"#'{prev_month}'!D2"
        pc.font = Font(name=FONT_BODY, size=9, color=COLORS['accent'], underline='single')
        pc.fill = PatternFill(start_color=COLORS['bg'], end_color=COLORS['bg'], fill_type='solid')
        pc.alignment = Alignment(horizontal='right', vertical='center')

    if next_month:
        nc = ws.cell(row=2, column=15)
        nc.value = f"{next_month} →"
        nc.hyperlink = f"#'{next_month}'!D2"
        nc.font = Font(name=FONT_BODY, size=9, color=COLORS['accent'], underline='single')
        nc.fill = PatternFill(start_color=COLORS['bg'], end_color=COLORS['bg'], fill_type='solid')
        nc.alignment = Alignment(horizontal='left', vertical='center')

    # ── INCOME SECTION ────────────────────────────────────────
    add_gold_rule(ws, row=6, start_col=4, end_col=9)
    _sec(ws, _INCOME_HDR_ROW, 4, "INCOME", end_col=9)
    _col_hdr(ws, _INCOME_COL_ROW, 4,
             ["Source", "Expected", "Received", "Difference", "% of Goal", ""])

    income_sample = {
        "January": [("Primary Salary", 4500, 4500),
                    ("Freelance Work", 800, 950),
                    ("Other Income", 0, 0)],
    }
    sample_rows = income_sample.get(month_name, [("Primary Salary", 4500, 0),
                                                  ("Freelance Work", 800, 0),
                                                  ("Other Income", 0, 0)])

    for i in range(_INCOME_DATA_END - _INCOME_DATA_START + 1):
        row = _INCOME_DATA_START + i
        alt = (i % 2 == 1)
        if i < len(sample_rows):
            source, expected, received = sample_rows[i]
        else:
            source, expected, received = "", "", ""

        diff_f   = f"=F{row}-E{row}" if source else ""
        pct_f    = f"=IF(E{row}>0,F{row}/E{row},\"\")" if source else ""

        _data(ws, row, 4, [source, expected if expected != "" else "",
                            received if received != "" else "",
                            diff_f, pct_f, ""],
              alt=alt,
              fmt=[None, '"$"#,##0.00', '"$"#,##0.00', '"$"#,##0.00', '0%', None])

        if source:
            add_dropdown(ws, f'D{row}', INCOME_CATS)

    # Income total row
    e_sum = f"=SUM(E{_INCOME_DATA_START}:E{_INCOME_DATA_END})"
    f_sum = f"=SUM(F{_INCOME_DATA_START}:F{_INCOME_DATA_END})"
    _total(ws, _INCOME_TOTAL_ROW, 4,
           ["TOTAL INCOME", e_sum, f_sum,
            f"=F{_INCOME_TOTAL_ROW}-E{_INCOME_TOTAL_ROW}", "", ""])
    _pct(ws, _INCOME_TOTAL_ROW, 8)

    # ── EXPENSE SECTION ───────────────────────────────────────
    add_gold_rule(ws, row=_EXPENSE_HDR_ROW - 1, start_col=4, end_col=9)
    _sec(ws, _EXPENSE_HDR_ROW, 4, "EXPENSES", end_col=9)
    _col_hdr(ws, _EXPENSE_COL_ROW, 4,
             ["Category", "Budgeted", "Actual", "Remaining", "% Used", "Status"])

    # Default budget amounts (realistic sample)
    default_budgets = {
        "Housing": 1800, "Utilities": 150, "Groceries": 500,
        "Dining Out": 200, "Transportation": 300, "Healthcare": 100,
        "Insurance": 280, "Personal Care": 80, "Clothing": 60,
        "Entertainment": 80, "Subscriptions": 50, "Education": 40,
        "Gifts & Donations": 60, "Savings": 500, "Debt Payment": 385,
        "Home Maintenance": 75, "Travel": 100, "Pets": 60,
        "Miscellaneous": 50,
    }

    for i, cat in enumerate(EXPENSE_CATS):
        row = _EXPENSE_DATA_START + i
        alt = (i % 2 == 1)
        budget = default_budgets.get(cat, 0)

        # SUMIFS from Variable Transactions: Category match + Month match
        actual_f = (f"=SUMIFS('Variable Transactions'!$G:$G,"
                    f"'Variable Transactions'!$F:$F,$D{row},"
                    f"'Variable Transactions'!$I:$I,\"{month_name}\")")

        remaining_f = f"=E{row}-F{row}"
        pct_f       = f"=IF(E{row}>0,F{row}/E{row},0)"
        status_f    = (f"=IF(E{row}=0,\"\","
                       f"IF(G{row}<0,\"Over Budget\","
                       f"IF(H{row}>0.9,\"Near Limit\",\"On Track\")))")

        _data(ws, row, 4,
              [cat, budget, actual_f, remaining_f, pct_f, status_f],
              alt=alt,
              fmt=[None, '"$"#,##0.00', '"$"#,##0.00', '"$"#,##0.00', '0%', None])

    # Expense total row
    e_exp = f"=SUM(E{_EXPENSE_DATA_START}:E{_EXPENSE_DATA_END})"
    f_exp = f"=SUM(F{_EXPENSE_DATA_START}:F{_EXPENSE_DATA_END})"
    _total(ws, _EXPENSE_TOTAL_ROW, 4,
           ["TOTAL EXPENSES", e_exp, f_exp,
            f"=E{_EXPENSE_TOTAL_ROW}-F{_EXPENSE_TOTAL_ROW}", "", ""])

    # Conditional formatting on Remaining (col G) for expense rows
    _add_over_budget_cf(ws, _EXPENSE_DATA_START, _EXPENSE_DATA_END)

    # Data bars on Actual (col F) expenses
    add_data_bars(ws,
                  f'F{_EXPENSE_DATA_START}:F{_EXPENSE_DATA_END}',
                  color=COLORS['accent'])

    # ── MONTHLY SUMMARY ───────────────────────────────────────
    add_gold_rule(ws, row=_SUMMARY_HDR_ROW - 1, start_col=4, end_col=9)
    _sec(ws, _SUMMARY_HDR_ROW, 4, "MONTHLY SUMMARY", end_col=9)
    _col_hdr(ws, _SUMMARY_COL_ROW, 4,
             ["", "Budgeted", "Actual", "Difference", "", ""])

    r42 = _SUMMARY_DATA_ROW
    r43 = r42 + 1
    r44 = r42 + 2
    r45 = r42 + 3

    summary_rows = [
        (r42, "Total Income",
         f"=E{_INCOME_TOTAL_ROW}", f"=F{_INCOME_TOTAL_ROW}"),
        (r43, "Total Expenses",
         f"=E{_EXPENSE_TOTAL_ROW}", f"=F{_EXPENSE_TOTAL_ROW}"),
        (r44, "Net Savings",
         f"=E{r42}-E{r43}", f"=F{r42}-F{r43}"),
    ]

    for row, label, budgeted, actual in summary_rows:
        alt = (row % 2 == 0)
        diff = f"=F{row}-E{row}"
        _data(ws, row, 4, [label, budgeted, actual, diff, "", ""],
              alt=alt,
              fmt=[None, '"$"#,##0.00', '"$"#,##0.00', '"$"#,##0.00', None, None])

    # Savings Rate row
    _data(ws, r45, 4,
          ["Savings Rate",
           f"=IF(E{r42}>0,(E{r42}-E{r43})/E{r42},0)",
           f"=IF(F{r42}>0,(F{r42}-F{r43})/F{r42},0)",
           "", "", ""],
          alt=(r45 % 2 == 0),
          fmt=[None, '0.0%', '0.0%', None, None, None])

    # ── PRINT-READY SUMMARY AREA ──────────────────────────────
    # (below main content — clean layout for printing)
    pr = r45 + 4
    add_gold_rule(ws, row=pr, start_col=4, end_col=9)
    note = ws.cell(row=pr + 1, column=4)
    note.value = (f"  {month_name} summary auto-calculated from Variable Transactions. "
                  "Enter Budget (col E) and income Received (col F) manually. "
                  "Expense Actual (col F) updates automatically as you log spending.")
    note.style = 'note_cell'
    ws.merge_cells(start_row=pr + 1, start_column=4, end_row=pr + 2, end_column=9)
    ws.row_dimensions[pr + 1].height = 28
