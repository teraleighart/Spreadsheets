from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule, DataBarRule
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


# ============================================================
# TERA LEIGH CREATIVE — BRAND TOKENS
# ============================================================

COLORS = {
    # Core brand colors
    'primary': '3D2535',      # Deep Plum — authority, structure, titles
    'accent': 'C4837A',       # Dusty Rose — action / CTA
    'blush': 'E8B4AD',        # Blush — softness / supporting emphasis
    'gold': 'B8914A',         # Antique Gold — one prestige moment per sheet
    'success': '7A8C6E',      # Muted Sage — semantic success/growth only
    'danger': 'C4522A',       # Terra Rust — alert / overdue / warning
    'bg': 'FAF5EE',           # Aged Cream — main workbook background
    'sidebar': 'F5EEE6',      # Warm cream variant for sidebar / nav
    'data': 'FFFDF7',         # Data cell fill — soft warm off-cream
    'data_alt': 'F8F1EA',     # Alternate data row — warm neutral, NOT white
    'border': 'E0D5CC',       # Warm rule / border
    'muted_text': '8A7F79',   # Warm muted text
    'body_text': '3D2535',    # Plum body text for stronger brand consistency
    'white': 'FFFFFF',
}


# ============================================================
# FONT FALLBACKS
# ============================================================
# Use brand fonts when available. If user machine does not have them,
# Excel will fall back automatically. These are still the correct choices
# to specify in the workbook.

FONT_DISPLAY = 'Cormorant Garamond'
FONT_BODY = 'Jost'
FONT_MONO = 'DM Mono'


# ============================================================
# NAMED STYLES
# ============================================================

def safe_add_named_style(wb, style):
    existing = {ns.name for ns in wb._named_styles}
    if style.name not in existing:
        wb.add_named_style(style)


def create_styles(wb, colors=COLORS):
    styles = {}

    styles['title'] = NamedStyle(
        name='title',
        font=Font(name=FONT_DISPLAY, size=20, bold=True, color=colors['primary']),
        alignment=Alignment(vertical='center', horizontal='left')
    )

    styles['subtitle'] = NamedStyle(
        name='subtitle',
        font=Font(name=FONT_BODY, size=11, color=colors['muted_text']),
        alignment=Alignment(vertical='center', horizontal='left')
    )

    # Structural headers should be plum, not sage
    styles['section_header'] = NamedStyle(
        name='section_header',
        font=Font(name=FONT_BODY, size=11, bold=True, color=colors['white']),
        fill=PatternFill(start_color=colors['primary'], end_color=colors['primary'], fill_type='solid'),
        alignment=Alignment(horizontal='left', vertical='center'),
        border=Border(bottom=Side(style='thin', color=colors['gold']))
    )

    styles['column_header'] = NamedStyle(
        name='column_header',
        font=Font(name=FONT_BODY, size=10, bold=True, color=colors['primary']),
        fill=PatternFill(start_color=colors['blush'], end_color=colors['blush'], fill_type='solid'),
        alignment=Alignment(horizontal='center', vertical='center', wrap_text=True),
        border=Border(
            top=Side(style='thin', color=colors['border']),
            bottom=Side(style='thin', color=colors['primary'])
        )
    )

    styles['data_cell'] = NamedStyle(
        name='data_cell',
        font=Font(name=FONT_BODY, size=10, color=colors['body_text']),
        fill=PatternFill(start_color=colors['data'], end_color=colors['data'], fill_type='solid'),
        alignment=Alignment(vertical='center', horizontal='left'),
        border=Border(bottom=Side(style='thin', color=colors['border']))
    )

    styles['data_alt'] = NamedStyle(
        name='data_alt',
        font=Font(name=FONT_BODY, size=10, color=colors['body_text']),
        fill=PatternFill(start_color=colors['data_alt'], end_color=colors['data_alt'], fill_type='solid'),
        alignment=Alignment(vertical='center', horizontal='left'),
        border=Border(bottom=Side(style='thin', color=colors['border']))
    )

    styles['note_cell'] = NamedStyle(
        name='note_cell',
        font=Font(name=FONT_BODY, size=10, italic=True, color=colors['muted_text']),
        fill=PatternFill(start_color=colors['bg'], end_color=colors['bg'], fill_type='solid'),
        alignment=Alignment(vertical='top', horizontal='left', wrap_text=True)
    )

    styles['kpi_value'] = NamedStyle(
        name='kpi_value',
        font=Font(name=FONT_DISPLAY, size=18, bold=True, color=colors['primary']),
        alignment=Alignment(horizontal='center', vertical='center')
    )

    styles['kpi_label'] = NamedStyle(
        name='kpi_label',
        font=Font(name=FONT_BODY, size=9, color=colors['muted_text']),
        alignment=Alignment(horizontal='center', vertical='center', wrap_text=True)
    )

    styles['total_row'] = NamedStyle(
        name='total_row',
        font=Font(name=FONT_BODY, size=10, bold=True, color=colors['primary']),
        fill=PatternFill(start_color='F3E8E2', end_color='F3E8E2', fill_type='solid'),
        alignment=Alignment(vertical='center'),
        border=Border(
            top=Side(style='thin', color=colors['primary']),
            bottom=Side(style='double', color=colors['primary'])
        )
    )

    styles['sidebar_link'] = NamedStyle(
        name='sidebar_link',
        font=Font(name=FONT_BODY, size=10, color=colors['primary'], underline='single'),
        fill=PatternFill(start_color=colors['sidebar'], end_color=colors['sidebar'], fill_type='solid'),
        alignment=Alignment(vertical='center', horizontal='left')
    )

    # Active navigation should feel authoritative, so plum instead of sage
    styles['sidebar_active'] = NamedStyle(
        name='sidebar_active',
        font=Font(name=FONT_BODY, size=10, bold=True, color=colors['white'], underline='single'),
        fill=PatternFill(start_color=colors['primary'], end_color=colors['primary'], fill_type='solid'),
        alignment=Alignment(vertical='center', horizontal='left'),
        border=Border(left=Side(style='medium', color=colors['gold']))
    )

    styles['eyebrow'] = NamedStyle(
        name='eyebrow',
        font=Font(name=FONT_MONO, size=9, color=colors['gold']),
        alignment=Alignment(vertical='center', horizontal='left')
    )

    styles['cta_cell'] = NamedStyle(
        name='cta_cell',
        font=Font(name=FONT_BODY, size=10, bold=True, color=colors['white']),
        fill=PatternFill(start_color=colors['accent'], end_color=colors['accent'], fill_type='solid'),
        alignment=Alignment(vertical='center', horizontal='center')
    )

    for style in styles.values():
        safe_add_named_style(wb, style)

    return styles


# ============================================================
# SHEET SETUP
# ============================================================

def setup_sheet(ws, title, colors=COLORS, active_tab=None, subtitle=None, sections=None):
    # Warm cream visible canvas
    bg = PatternFill(start_color=colors['bg'], end_color=colors['bg'], fill_type='solid')
    for row in ws.iter_rows(min_row=1, max_row=120, min_col=1, max_col=50):
        for cell in row:
            cell.fill = bg

    # Gridlines OFF — aligns better with calm editorial look
    ws.sheet_view.showGridLines = False

    # Tab color should be plum for consistency
    ws.sheet_properties.tabColor = colors['primary']

    # Row heights
    for r in range(1, 121):
        ws.row_dimensions[r].height = 18
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 20

    # Sidebar / gutter columns
    ws.column_dimensions['A'].width = 3.25
    ws.column_dimensions['B'].width = 18.5
    ws.column_dimensions['C'].width = 3.25

    # Main content columns can be tuned later per sheet
    ws.freeze_panes = 'D5'

    # Title
    ws['D2'] = title
    ws['D2'].style = 'title'

    # Optional subtitle / plain-English orienting sentence
    if subtitle:
        ws['D3'] = subtitle
        ws['D3'].style = 'subtitle'

    # One subtle gold divider = editorial signature
    add_gold_rule(ws, row=4, start_col=4, end_col=14, colors=colors)

    # Sidebar
    add_sidebar(ws, colors=colors, active_tab=active_tab or ws.title, sections=sections)


def add_gold_rule(ws, row, start_col, end_col, colors=COLORS):
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = Border(top=Side(style='thin', color=colors['gold']))


# ============================================================
# SIDEBAR / NAVIGATION
# ============================================================

def add_sidebar(ws, colors=COLORS, active_tab=None, sections=None):
    sidebar_fill = PatternFill(start_color=colors['sidebar'], end_color=colors['sidebar'], fill_type='solid')

    for r in range(1, 70):
        ws.cell(row=r, column=1).fill = sidebar_fill
        ws.cell(row=r, column=2).fill = sidebar_fill

    # Optional countdown / helper area
    ws['B2'] = 'START HERE'
    ws['B2'].style = 'eyebrow'

    if sections is None:
        sections = {
            7: ("OVERVIEW", None),
            8: (None, "Instructions"),
            9: (None, "Setup"),
            10: (None, "Dashboard"),
        }

    for row, (section, link) in sections.items():
        if section:
            label_cell = ws.cell(row=row, column=1)
            label_cell.value = section
            label_cell.font = Font(name=FONT_BODY, size=8, bold=True, color=colors['muted_text'])
            label_cell.fill = sidebar_fill

        if link:
            cell = ws.cell(row=row, column=2)
            cell.value = link
            cell.hyperlink = f"#'{link}'!D2"

            if link == active_tab:
                cell.style = 'sidebar_active'
            else:
                cell.style = 'sidebar_link'


# ============================================================
# KPI CARDS
# ============================================================

def add_kpi_card(
    ws,
    start_row,
    start_col,
    label,
    value_formula,
    fmt='#,##0',
    colors=COLORS,
    accent_color=None
):
    card_fill = PatternFill(start_color=colors['data'], end_color=colors['data'], fill_type='solid')
    border = Border(
        left=Side(style='thin', color=colors['border']),
        right=Side(style='thin', color=colors['border']),
        top=Side(style='thin', color=colors['border']),
        bottom=Side(style='thin', color=colors['border'])
    )

    # Plum by default for structural cards
    top_accent = accent_color or colors['primary']

    ws.merge_cells(
        start_row=start_row, start_column=start_col,
        end_row=start_row + 2, end_column=start_col + 2
    )

    # Apply fill and borders first
    for r in range(start_row, start_row + 3):
        for c in range(start_col, start_col + 3):
            ws.cell(row=r, column=c).fill = card_fill
            ws.cell(row=r, column=c).border = border

    # Top accent strip
    for c in range(start_col, start_col + 3):
        cell = ws.cell(row=start_row, column=c)
        cell.border = Border(
            top=Side(style='medium', color=top_accent),
            left=Side(style='thin', color=colors['border']),
            right=Side(style='thin', color=colors['border']),
            bottom=Side(style='thin', color=colors['border'])
        )

    # Value
    val_cell = ws.cell(row=start_row + 1, column=start_col)
    val_cell.value = value_formula
    val_cell.style = 'kpi_value'
    val_cell.number_format = fmt
    val_cell.fill = card_fill

    # Label
    lbl_cell = ws.cell(row=start_row + 2, column=start_col)
    lbl_cell.value = label
    lbl_cell.style = 'kpi_label'
    lbl_cell.fill = card_fill


# ============================================================
# CHART STYLING
# ============================================================

def style_chart(chart, title, chart_type='bar', colors=COLORS, width=15, height=9):
    chart.title = title
    chart.style = 10
    chart.width = width
    chart.height = height

    chart.dataLabels = DataLabelList()
    if chart_type in ('donut', 'pie'):
        chart.dataLabels.showPercent = True
        chart.dataLabels.showCatName = True
        chart.dataLabels.showVal = False
    else:
        chart.dataLabels.showVal = True

    if chart.legend:
        chart.legend.position = 'b'

    # Cleaner chart frame
    chart.plot_area.graphicalProperties = None

    return chart


# ============================================================
# DATA TABLES
# ============================================================

def add_data_table(ws, headers, data_rows, start_row, start_col, colors=COLORS):
    # Headers
    for i, header in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + i)
        cell.value = header
        cell.style = 'column_header'

    # Rows
    for r_idx, row_data in enumerate(data_rows):
        style_name = 'data_cell' if r_idx % 2 == 0 else 'data_alt'
        for c_idx, value in enumerate(row_data):
            cell = ws.cell(row=start_row + 1 + r_idx, column=start_col + c_idx)
            cell.value = value
            cell.style = style_name

            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # Gentle gold rule above the table block
    add_gold_rule(ws, row=start_row - 1, start_col=start_col, end_col=start_col + len(headers) - 1, colors=colors)


# ============================================================
# CONDITIONAL FORMATTING
# ============================================================

def add_status_formatting(ws, cell_range, colors=COLORS):
    complete_style = DifferentialStyle(
        fill=PatternFill(start_color=colors['success'], end_color=colors['success'], fill_type='solid')
    )
    progress_style = DifferentialStyle(
        fill=PatternFill(start_color=colors['accent'], end_color=colors['accent'], fill_type='solid')
    )
    overdue_style = DifferentialStyle(
        fill=PatternFill(start_color=colors['danger'], end_color=colors['danger'], fill_type='solid')
    )

    first_cell = cell_range.split(':')[0]

    ws.conditional_formatting.add(
        cell_range,
        Rule(
            type='containsText',
            operator='containsText',
            text='Complete',
            dxf=complete_style,
            formula=[f'NOT(ISERROR(SEARCH("Complete",{first_cell})))']
        )
    )

    ws.conditional_formatting.add(
        cell_range,
        Rule(
            type='containsText',
            operator='containsText',
            text='In Progress',
            dxf=progress_style,
            formula=[f'NOT(ISERROR(SEARCH("In Progress",{first_cell})))']
        )
    )

    ws.conditional_formatting.add(
        cell_range,
        Rule(
            type='containsText',
            operator='containsText',
            text='Overdue',
            dxf=overdue_style,
            formula=[f'NOT(ISERROR(SEARCH("Overdue",{first_cell})))']
        )
    )


def add_data_bars(ws, cell_range, color=None):
    # Plum by default for structural data visualization
    rule = DataBarRule(
        start_type='min',
        end_type='max',
        color=color or COLORS['primary']
    )
    ws.conditional_formatting.add(cell_range, rule)


# ============================================================
# PLAIN-ENGLISH SHEET TITLE HELPER
# ============================================================

def add_sheet_intro(ws, heading, body, row=6, col=4):
    ws.cell(row=row, column=col).value = heading
    ws.cell(row=row, column=col).style = 'section_header'

    ws.merge_cells(
        start_row=row + 1, start_column=col,
        end_row=row + 2, end_column=col + 5
    )

    body_cell = ws.cell(row=row + 1, column=col)
    body_cell.value = body
    body_cell.style = 'note_cell'


# ============================================================
# QUALITY GATE
# ============================================================

def run_quality_gate(wb):
    issues = []

    for ws in wb.worksheets:
        if ws.sheet_view.showGridLines:
            issues.append(f"{ws.title}: gridlines still ON")

        if not ws.sheet_properties.tabColor:
            issues.append(f"{ws.title}: no tab color")

        for c in ws._charts:
            if not c.title:
                issues.append(f"{ws.title}: chart without title")

        # Soft check for obvious brand drift
        if ws['D2'].font.name not in (FONT_DISPLAY, 'Georgia', None):
            issues.append(f"{ws.title}: title font drifted from brand display font")

    return issues


# ============================================================
# DATA VALIDATION HELPERS
# ============================================================

def add_dropdown(ws, cell_range, options, allow_blank=True):
    """Inline list dropdown — use for short, fixed option sets (≤ ~8 items)."""
    joined = ",".join(options)
    dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=allow_blank)
    ws.add_data_validation(dv)
    dv.add(cell_range)


def add_range_dropdown(ws, cell_range, formula_range, allow_blank=True):
    """Range-based dropdown — use when options live in a named range or sheet ref."""
    dv = DataValidation(type="list", formula1=formula_range, allow_blank=allow_blank)
    ws.add_data_validation(dv)
    dv.add(cell_range)


# ============================================================
# PRODUCT CATEGORY BLUEPRINTS
# ============================================================

BLUEPRINTS = {
    "Annual Budget Spreadsheet": {
        "tabs": [
            "Instructions", "Setup", "Bank Accounts", "Recurring Transactions", "Payments",
            "Variable Transactions", "All-in-One Dashboard", "Annual Totals", "Automated Calendar",
            "Paycheck Dashboard", "January", "February", "March", "April", "May", "June",
            "July", "August", "September", "October", "November", "December",
            "50/30/20 Dashboard", "Expense Distribution", "Sinking Funds", "Debt Calculator",
            "Net Worth", "Investment Forecast", "No-Spending Challenge",
        ],
        "features": [
            "Currency dropdowns from Setup",
            "SUMIF by category",
            "Monthly auto-totals",
            "Dashboard with 4+ charts",
            "Conditional formatting on over-budget items",
            "Data bars on amounts",
        ],
    },

    "Book Tracker": {
        "tabs": [
            "Instructions", "Setup", "Book Tracker", "Books Gallery",
            "Digital Bookshelf", "Reading Calendar", "Wishlist", "All-in-One Dashboard",
        ],
        "features": [
            "Genre/status/rating dropdowns",
            "COUNTIF and AVERAGEIF stats",
            "Reading streak tracking",
            "Goal progress",
            "Dashboard charts",
        ],
    },

    "Wedding Planner": {
        "tabs": [
            "Instructions", "Setup", "Save the Date", "Theme", "Dashboard", "Calendar",
            "Timeline", "Itinerary", "Packing List", "Vendors Choice", "Venue Options",
            "Budget", "Contact Info", "Guest List", "Seating Plan", "Wedding Party",
            "Food & Drinks", "Photoshoot", "Photo Gallery", "Music",
            "Gifts & Thank You", "Honeymoon",
        ],
        "features": [
            "Days-left countdown",
            "Cross-tab dashboard with 7 charts",
            "Budget tracking",
            "Guest list with RSVP and meal dropdowns",
            "Seating plan grid",
            "Vendor comparison",
            "Timeline with priority levels",
        ],
    },

    "Fitness Tracker": {
        "tabs": [
            "Instructions", "Setup", "Dashboard", "Workout Log", "Meal Planner",
            "Progress Photos", "Body Measurements", "Goals",
            "Weekly Summary", "Monthly Summary", "Exercise Library",
        ],
        "features": [
            "Workout tracking",
            "Meal planning",
            "Measurement logging",
            "Goals and summaries",
            "Dashboard charts",
        ],
    },

    "Project Manager": {
        "tabs": [
            "Instructions", "Setup", "Dashboard", "Tasks", "Timeline/Gantt",
            "Team Members", "Budget", "Notes", "Archive",
        ],
        "features": [
            "Task tracking",
            "Timeline/Gantt view",
            "Budget tracking",
            "Team assignments",
            "Dashboard overview",
        ],
    },
}
